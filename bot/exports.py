from __future__ import annotations
"""Генерация файлов экспорта (CSV/XLSX/PDF) из БД и отправка через Telegram.

Используется publisher-циклом: находит pending-задания в cd_exports,
генерирует файл, отправляет документом в Telegram пользователю, помечает done.
Отправка — синхронный urllib multipart (надёжно в потоке asyncio.to_thread,
в отличие от aiogram Bot + asyncio.run внутри потока).
"""
import csv
import io
import json
import logging
import os
import uuid
from datetime import datetime
from pathlib import Path
from urllib.request import Request, urlopen

import psycopg

from bot.db import db_url

logger = logging.getLogger('channeldesk.exports')

FONTS_DIR = Path(__file__).resolve().parents[1] / 'assets' / 'fonts'
TELEGRAM_API = 'https://api.telegram.org'


def _fmt(dt) -> str:
    if not dt:
        return ''
    if isinstance(dt, datetime):
        return dt.strftime('%Y-%m-%d %H:%M')
    return str(dt)


def _send_document(token: str, chat_id: int, filename: str, data: bytes, caption: str) -> None:
    """Отправляет документ через Telegram Bot API (multipart/form-data, синхронно)."""
    boundary = '----channeldesk' + uuid.uuid4().hex
    body = io.BytesIO()
    for name, value in (('chat_id', str(chat_id)), ('caption', caption)):
        body.write(f'--{boundary}\r\n'.encode())
        body.write(f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode())
        body.write(f'{value}\r\n'.encode())
    body.write(f'--{boundary}\r\n'.encode())
    body.write(f'Content-Disposition: form-data; name="document"; filename="{filename}"\r\n'.encode())
    body.write(b'Content-Type: application/octet-stream\r\n\r\n')
    body.write(data)
    body.write(f'\r\n--{boundary}--\r\n'.encode())
    req = Request(f'{TELEGRAM_API}/bot{token}/sendDocument', data=body.getvalue(), method='POST',
                  headers={'Content-Type': f'multipart/form-data; boundary={boundary}'})
    with urlopen(req, timeout=120) as resp:
        payload = json.loads(resp.read().decode('utf-8'))
    if not payload.get('ok'):
        raise RuntimeError(f"Telegram API error: {payload.get('description', 'unknown')}")


def _fmt(dt) -> str:
    if not dt:
        return ''
    if isinstance(dt, datetime):
        return dt.strftime('%Y-%m-%d %H:%M')
    return str(dt)


def _load_rows(conn, kind: str, workspace_id: int, period_year: int | None = None,
               period_month: int | None = None) -> list[dict]:
    with conn.cursor() as cur:
        if kind == 'posts':
            cur.execute("""SELECT p.*, c.title AS channel_title FROM cd_posts p
            LEFT JOIN cd_channels c ON c.id=p.channel_id
            WHERE p.workspace_id=%s ORDER BY p.updated_at DESC""", (workspace_id,))
        elif kind == 'bookings':
            cur.execute("""SELECT b.*, a.name AS advertiser_name, c.title AS channel_title
            FROM cd_ad_bookings b LEFT JOIN cd_advertisers a ON a.id=b.advertiser_id
            LEFT JOIN cd_channels c ON c.id=b.channel_id
            WHERE b.workspace_id=%s ORDER BY b.id DESC""", (workspace_id,))
        elif kind == 'media_kits':
            cur.execute("""SELECT mk.*, c.title AS channel_title
            FROM cd_media_kits mk LEFT JOIN cd_channels c ON c.id=mk.channel_id
            WHERE mk.workspace_id=%s AND mk.is_active=true ORDER BY mk.name""", (workspace_id,))
        elif period_year is not None and period_month is not None:
            cur.execute("""SELECT * FROM cd_finance_transactions WHERE workspace_id=%s
            AND occurred_at >= make_date(%s,%s,1)
            AND occurred_at < make_date(%s,%s,1) + interval '1 month'
            ORDER BY occurred_at DESC, id DESC""",
                        (workspace_id, period_year, period_month, period_year, period_month))
        else:
            cur.execute("""SELECT * FROM cd_finance_transactions WHERE workspace_id=%s
            ORDER BY occurred_at DESC, id DESC""", (workspace_id,))
        return cur.fetchall() or []


def _csv_bytes(rows: list[dict], kind: str) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=';')
    if kind == 'posts':
        writer.writerow(['ID', 'Заголовок', 'Текст', 'Статус', 'Канал', 'Запланировано', 'Опубликовано', 'Ошибка'])
        for p in rows:
            writer.writerow([p['id'], p.get('title') or '', (p.get('text') or '')[:200], p.get('status') or '',
                             p.get('channel_title') or '', _fmt(p.get('scheduled_at')), _fmt(p.get('published_at')),
                             p.get('last_error') or ''])
    elif kind == 'bookings':
        writer.writerow(['ID', 'Рекламодатель', 'Канал', 'Формат', 'Стоимость', 'Валюта', 'Статус', 'Оплата',
                         'Публикация', 'ERID', 'ERID требуется'])
        for b in rows:
            writer.writerow([b['id'], b.get('advertiser_name') or '', b.get('channel_title') or '',
                             b.get('format') or '', b.get('cost'), b.get('currency') or '', b.get('status') or '',
                             b.get('payment_status') or '', _fmt(b.get('publish_at')), b.get('erid') or '',
                             'да' if b.get('erid_required', True) else 'нет'])
    else:
        writer.writerow(['ID', 'Тип', 'Сумма', 'Валюта', 'Категория', 'Описание', 'Дата'])
        for t in rows:
            writer.writerow([t['id'], 'Доход' if t.get('type') == 'income' else 'Расход', t.get('amount'),
                             t.get('currency') or '', t.get('category') or '', t.get('description') or '',
                             _fmt(t.get('occurred_at'))])
    return ('\ufeff' + buffer.getvalue()).encode('utf-8')


def _xlsx_bytes(rows: list[dict], kind: str) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    if kind == 'posts':
        ws.title = 'Публикации'
        ws.append(['ID', 'Заголовок', 'Текст', 'Статус', 'Канал', 'Запланировано', 'Опубликовано', 'Ошибка'])
        for p in rows:
            ws.append([p['id'], p.get('title') or '', (p.get('text') or '')[:200], p.get('status') or '',
                       p.get('channel_title') or '', _fmt(p.get('scheduled_at')), _fmt(p.get('published_at')),
                       p.get('last_error') or ''])
    elif kind == 'bookings':
        ws.title = 'Брони'
        ws.append(['ID', 'Рекламодатель', 'Канал', 'Формат', 'Стоимость', 'Валюта', 'Статус', 'Оплата',
                   'Публикация', 'ERID', 'ERID требуется'])
        for b in rows:
            ws.append([b['id'], b.get('advertiser_name') or '', b.get('channel_title') or '', b.get('format') or '',
                       float(b.get('cost') or 0), b.get('currency') or '', b.get('status') or '',
                       b.get('payment_status') or '', _fmt(b.get('publish_at')), b.get('erid') or '',
                       'да' if b.get('erid_required', True) else 'нет'])
    else:
        ws.title = 'Финансы'
        ws.append(['ID', 'Тип', 'Сумма', 'Валюта', 'Категория', 'Описание', 'Дата'])
        for t in rows:
            ws.append([t['id'], 'Доход' if t.get('type') == 'income' else 'Расход', float(t.get('amount') or 0),
                       t.get('currency') or '', t.get('category') or '', t.get('description') or '',
                       _fmt(t.get('occurred_at'))])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _pdf_bytes(rows: list[dict], kind: str = 'posts') -> bytes:
    """Собирает PDF именно для выбранного типа экспорта, а не один вечный PDF постов."""
    from fpdf import FPDF

    pdf = FPDF(orientation='L' if kind in ('bookings', 'finance') else 'P')
    pdf.add_font('DejaVu', '', str(FONTS_DIR / 'DejaVuSans.ttf'))
    pdf.add_font('DejaVu', 'B', str(FONTS_DIR / 'DejaVuSans-Bold.ttf'))
    pdf.add_page()
    if kind == 'media_kits':
        for index, kit in enumerate(rows):
            if index:
                pdf.add_page()
            pdf.set_font('DejaVu', 'B', 18)
            pdf.cell(0, 11, kit.get('name') or 'Медиакит', new_x='LMARGIN', new_y='NEXT')
            pdf.set_font('DejaVu', '', 10)
            pdf.cell(0, 7, f"Канал: {kit.get('channel_title') or 'не указан'}",
                     new_x='LMARGIN', new_y='NEXT')
            pdf.ln(3)
            if kit.get('description'):
                pdf.set_font('DejaVu', 'B', 11)
                pdf.cell(0, 7, 'О канале', new_x='LMARGIN', new_y='NEXT')
                pdf.set_font('DejaVu', '', 10)
                pdf.multi_cell(0, 6, str(kit['description']), new_x='LMARGIN', new_y='NEXT')
            stats = kit.get('stats') or {}
            if stats:
                pdf.set_font('DejaVu', 'B', 11)
                pdf.cell(0, 7, 'Статистика', new_x='LMARGIN', new_y='NEXT')
                pdf.set_font('DejaVu', '', 10)
                for key, value in stats.items():
                    pdf.cell(0, 6, f'{key}: {value}', new_x='LMARGIN', new_y='NEXT')
            pricing = kit.get('pricing') or []
            if pricing:
                pdf.set_font('DejaVu', 'B', 11)
                pdf.cell(0, 7, 'Стоимость размещений', new_x='LMARGIN', new_y='NEXT')
                pdf.set_font('DejaVu', '', 10)
                for item in pricing:
                    if isinstance(item, dict):
                        fmt = item.get('format') or 'размещение'
                        price = item.get('price', '')
                        currency = item.get('currency') or 'RUB'
                        pdf.cell(0, 6, f'{fmt}: {price} {currency}', new_x='LMARGIN', new_y='NEXT')
                    else:
                        pdf.cell(0, 6, str(item), new_x='LMARGIN', new_y='NEXT')
            contacts = kit.get('contacts') or {}
            if contacts:
                pdf.set_font('DejaVu', 'B', 11)
                pdf.cell(0, 7, 'Контакты', new_x='LMARGIN', new_y='NEXT')
                pdf.set_font('DejaVu', '', 10)
                for key, value in contacts.items():
                    pdf.cell(0, 6, f'{key}: {value}', new_x='LMARGIN', new_y='NEXT')
        if not rows:
            pdf.set_font('DejaVu', 'B', 16)
            pdf.cell(0, 10, 'ChannelDesk - Медиакиты', new_x='LMARGIN', new_y='NEXT')
            pdf.set_font('DejaVu', '', 10)
            pdf.cell(0, 7, 'Медиакитов пока нет.', new_x='LMARGIN', new_y='NEXT')
        return pdf.output()

    pdf.set_font('DejaVu', 'B', 14)
    titles = {'posts': 'Публикации', 'bookings': 'Брони', 'finance': 'Финансы'}
    pdf.cell(0, 10, f'ChannelDesk - {titles.get(kind, kind)}', new_x='LMARGIN', new_y='NEXT')
    pdf.set_font('DejaVu', 'B', 8)

    if kind == 'finance':
        headers = ['ID', 'Тип', 'Сумма', 'Валюта', 'Категория', 'Описание', 'Дата']
        widths = [12, 24, 30, 18, 34, 116, 38]
        values = lambda row: [row.get('id'), 'Доход' if row.get('type') == 'income' else 'Расход',
                              row.get('amount'), row.get('currency') or '', row.get('category') or '',
                              row.get('description') or '', _fmt(row.get('occurred_at'))]
    elif kind == 'bookings':
        headers = ['ID', 'Рекламодатель', 'Канал', 'Формат', 'Стоимость', 'Статус', 'Оплата', 'Публикация']
        widths = [12, 62, 50, 28, 30, 30, 30, 45]
        values = lambda row: [row.get('id'), row.get('advertiser_name') or '', row.get('channel_title') or '',
                              row.get('format') or '', row.get('cost'), row.get('status') or '',
                              row.get('payment_status') or '', _fmt(row.get('publish_at'))]
    else:
        headers = ['ID', 'Заголовок', 'Статус', 'Канал', 'Запланировано']
        widths = [12, 68, 34, 45, 48]
        values = lambda row: [row.get('id'), row.get('title') or '', row.get('status') or '',
                              row.get('channel_title') or '', _fmt(row.get('scheduled_at'))]

    for header, width in zip(headers, widths):
        pdf.cell(width, 7, header, border=1)
    pdf.ln()
    pdf.set_font('DejaVu', '', 8)
    for row in rows:
        for value, width in zip(values(row), widths):
            text = '' if value is None else str(value)
            pdf.cell(width, 7, text[:70], border=1)
        pdf.ln()
    return pdf.output()


def _generate_file(conn, kind: str, fmt: str, workspace_id: int, period_year: int | None = None,
                   period_month: int | None = None) -> tuple[bytes, str, str]:
    rows = _load_rows(conn, kind, workspace_id, period_year, period_month)
    filename = f'{kind}.{fmt}'
    if fmt == 'csv':
        return _csv_bytes(rows, kind), filename, 'text/csv; charset=utf-8'
    if fmt == 'xlsx':
        return _xlsx_bytes(rows, kind), filename, \
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return _pdf_bytes(rows, kind), filename, 'application/pdf'


def _claim_export(conn) -> dict | None:
    """Забирает одно pending-задание. Простой SELECT (publisher один, гонок нет)."""
    with conn.cursor() as cur:
        cur.execute("""SELECT id,workspace_id,telegram_id,kind,format,period_year,period_month FROM cd_exports
        WHERE status='pending' ORDER BY created_at LIMIT 1""")
        job = cur.fetchone()
        if job:
            cur.execute("UPDATE cd_exports SET status='processing' WHERE id=%s", (job['id'],))
    return job


def process_pending_exports(token: str, conn) -> int:
    """Обрабатывает до 5 pending-заданий экспорта.

    ВАЖНО: работает на ОТДЕЛЬНОМ соединении с autocommit=True — полностью
    изолировано от длинной транзакции publisher-цикла. Через Supabase pooler
    общая транзакция может не видеть строки, записанные Vercel.
    """
    processed = 0
    try:
        work_conn = psycopg.connect(db_url(), row_factory=psycopg.rows.dict_row, autocommit=True)
    except Exception as exc:
        logger.exception('cannot open export connection: %s', exc)
        return 0
    try:
        for _ in range(5):
            job = _claim_export(work_conn)
            if not job:
                break
            try:
                data, filename, mime = _generate_file(
                    work_conn, job['kind'], job['format'], job['workspace_id'],
                    job.get('period_year'), job.get('period_month'))
                _send_document(token, job['telegram_id'], filename, data,
                               f'ChannelDesk: экспорт «{job["kind"]}» ({job["format"]})')
                with work_conn.cursor() as cur:
                    cur.execute("UPDATE cd_exports SET status='done',completed_at=now() WHERE id=%s", (job['id'],))
                processed += 1
                logger.info('export %s (%s.%s) sent to %s', job['id'], job['kind'], job['format'], job['telegram_id'])
            except Exception as exc:
                logger.exception('export %s failed', job['id'])
                with work_conn.cursor() as cur:
                    cur.execute("UPDATE cd_exports SET status='failed',error_text=%s WHERE id=%s",
                                (str(exc)[:500], job['id']))
        return processed
    finally:
        work_conn.close()
